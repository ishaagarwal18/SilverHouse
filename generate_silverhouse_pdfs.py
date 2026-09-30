import os
import re
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, KeepTogether, HRFlowable
)
from reportlab.pdfgen import canvas
import docx

class NumberedCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_number(num_pages)
            super().showPage()
        super().save()

    def draw_page_number(self, page_count):
        self.saveState()
        self.setFont("Helvetica", 8)
        self.setFillColor(colors.HexColor('#666666'))
        
        # Top running rule (pages 2+)
        if self._pageNumber > 1:
            self.setStrokeColor(colors.HexColor('#D4AF37'))
            self.setLineWidth(0.6)
            self.line(40, 800, 555, 800)
            self.drawString(40, 805, "SILVER HOUSE • Official Knowledge & Assurance Series")

        # Bottom footer
        self.setStrokeColor(colors.HexColor('#E2D7BE'))
        self.setLineWidth(0.6)
        self.line(40, 42, 555, 42)
        
        footer_text = "Silver House • 217, Kanak Chamber, Gandhi Road, Ahmedabad - 380058 • Mobile: +91 95371 78477 • Email: sunilag28017@gmail.com"
        self.drawString(40, 30, footer_text)
        
        page_str = f"Page {self._pageNumber} of {page_count}"
        self.drawRightString(555, 30, page_str)
        self.restoreState()

def clean_text(text):
    if not text:
        return ""
    # Remove citations like cite...
    text = re.sub(r'[\ue200-\ue2ff]|[\ue000-\uf8ff]', '', text)
    text = re.sub(r'cite.*?', '', text)
    text = re.sub(r'\[cite.*?\]', '', text)
    text = text.replace('•', '&bull;').replace('×', '&times;').replace('÷', '&divide;')
    text = text.replace('“', '&ldquo;').replace('”', '&rdquo;').replace('’', '&rsquo;').replace('‘', '&lsquo;')
    return text.strip()

def build_about_us_pdf(docx_path, output_pdf, logo_path):
    doc = SimpleDocTemplate(
        output_pdf,
        pagesize=A4,
        leftMargin=40,
        rightMargin=40,
        topMargin=48,
        bottomMargin=54
    )
    
    styles = getSampleStyleSheet()
    
    # Custom styles
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=24,
        leading=28,
        textColor=colors.HexColor('#1A1A1A'),
        spaceAfter=4
    )
    
    subtitle_style = ParagraphStyle(
        'DocSub',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=11,
        leading=15,
        textColor=colors.HexColor('#9E7D3B'),
        spaceAfter=14
    )
    
    h1_style = ParagraphStyle(
        'Heading1',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=13,
        leading=17,
        textColor=colors.HexColor('#1A1A1A'),
        spaceBefore=12,
        spaceAfter=6,
        keepWithNext=True
    )
    
    body_style = ParagraphStyle(
        'BodyDark',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9.5,
        leading=14.5,
        textColor=colors.HexColor('#2D3748'),
        spaceAfter=8
    )

    callout_style = ParagraphStyle(
        'CalloutText',
        parent=styles['Normal'],
        fontName='Helvetica-Oblique',
        fontSize=10,
        leading=15,
        textColor=colors.HexColor('#1A1A1A')
    )

    story = []

    # Header with Logo
    header_data = []
    if os.path.exists(logo_path):
        img = Image(logo_path, width=54, height=54)
        info_para = Paragraph(
            "<b>SILVER HOUSE</b><br/><font size='9' color='#9E7D3B'><b>Manufacturer & Trader of Silver Products</b></font><br/><font size='8' color='#666666'>217, Kanak Chamber, Gandhi Road, Ahmedabad - 380058</font>",
            body_style
        )
        header_table = Table([[img, info_para]], colWidths=[65, 450])
        header_table.setStyle(TableStyle([
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('BOTTOMPADDING', (0,0), (-1,-1), 6),
        ]))
        story.append(header_table)
        story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor('#D4AF37'), spaceBefore=4, spaceAfter=14))

    # Read docx
    d = docx.Document(docx_path)
    paras = [clean_text(p.text) for p in d.paragraphs if clean_text(p.text)]

    for idx, p in enumerate(paras):
        if idx == 0 and "SILVER HOUSE" in p:
            continue
        if idx == 1 and "About Us" in p:
            story.append(Paragraph("ABOUT SILVER HOUSE", title_style))
            story.append(Paragraph("From a Small Beginning in Ahmedabad to a Trusted Name in Silver", subtitle_style))
            continue
        
        # Section headers
        if p in ["Where It All Began", "Built Through Wholesale", "Growing Across Gujarat", "A New Chapter in Retail", "Silver House Today", "What We Believe"]:
            story.append(Spacer(1, 4))
            story.append(Paragraph(f"<b>{p}</b>", h1_style))
            story.append(HRFlowable(width="100%", thickness=0.8, color=colors.HexColor('#E2D7BE'), spaceBefore=2, spaceAfter=6))
        elif "Established in Ahmedabad in 2010" in p or "Strengthened through relationships" in p:
            # Highlight final signature block
            box_data = [[
                Paragraph(f"<b>{p}</b>", callout_style)
            ]]
            t = Table(box_data, colWidths=[515])
            t.setStyle(TableStyle([
                ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#F9F6EE')),
                ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#D4AF37')),
                ('TOPPADDING', (0,0), (-1,-1), 8),
                ('BOTTOMPADDING', (0,0), (-1,-1), 8),
                ('LEFTPADDING', (0,0), (-1,-1), 12),
                ('RIGHTPADDING', (0,0), (-1,-1), 12),
            ]))
            story.append(Spacer(1, 8))
            story.append(t)
        else:
            story.append(Paragraph(p, body_style))

    doc.build(story, canvasmaker=NumberedCanvas)
    print(f"[OK] Built {output_pdf} ({os.path.getsize(output_pdf)} bytes)")

def build_purity_guide_pdf(docx_path, output_pdf, logo_path):
    doc = SimpleDocTemplate(
        output_pdf,
        pagesize=A4,
        leftMargin=40,
        rightMargin=40,
        topMargin=48,
        bottomMargin=54
    )
    
    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=21,
        leading=25,
        textColor=colors.HexColor('#1A1A1A'),
        spaceAfter=4
    )
    
    subtitle_style = ParagraphStyle(
        'DocSub',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=11,
        leading=15,
        textColor=colors.HexColor('#9E7D3B'),
        spaceAfter=12
    )
    
    h1_style = ParagraphStyle(
        'Heading1',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=12,
        leading=16,
        textColor=colors.HexColor('#1A1A1A'),
        spaceBefore=10,
        spaceAfter=4,
        keepWithNext=True
    )
    
    body_style = ParagraphStyle(
        'BodyDark',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=13.5,
        textColor=colors.HexColor('#2D3748'),
        spaceAfter=6
    )

    box_style = ParagraphStyle(
        'BoxText',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=10,
        leading=14,
        textColor=colors.HexColor('#1A1A1A')
    )

    table_header_style = ParagraphStyle(
        'TH',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=9,
        leading=12,
        textColor=colors.white,
        alignment=1
    )

    table_cell_style = ParagraphStyle(
        'TD',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8.5,
        leading=11.5,
        textColor=colors.HexColor('#222222'),
        alignment=1
    )

    story = []

    # Header with Logo
    if os.path.exists(logo_path):
        img = Image(logo_path, width=54, height=54)
        info_para = Paragraph(
            "<b>SILVER HOUSE</b><br/><font size='9' color='#9E7D3B'><b>Official Silver Purity & Hallmarking Consumer Guide</b></font><br/><font size='8' color='#666666'>217, Kanak Chamber, Gandhi Road, Ahmedabad - 380058 | +91 95371 78477</font>",
            body_style
        )
        header_table = Table([[img, info_para]], colWidths=[65, 450])
        header_table.setStyle(TableStyle([
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('BOTTOMPADDING', (0,0), (-1,-1), 6),
        ]))
        story.append(header_table)
        story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor('#D4AF37'), spaceBefore=4, spaceAfter=10))

    story.append(Paragraph("THE SIMPLE GUIDE TO SILVER PURITY", title_style))
    story.append(Paragraph("Understanding 999 Fine Silver, 925 Sterling Silver & BIS Hallmarking Standards", subtitle_style))

    # Read docx
    d = docx.Document(docx_path)
    paras = [clean_text(p.text) for p in d.paragraphs if clean_text(p.text)]

    # Cheat sheet table data
    table_rows = [
        [Paragraph("<b>Purity Mark</b>", table_header_style), Paragraph("<b>Silver Content</b>", table_header_style), Paragraph("<b>Easy Way to Read It</b>", table_header_style)],
        [Paragraph("<b>999</b>", table_cell_style), Paragraph("99.9%", table_cell_style), Paragraph("999 parts pure silver out of 1,000", table_cell_style)],
        [Paragraph("<b>990</b>", table_cell_style), Paragraph("99.0%", table_cell_style), Paragraph("990 parts silver out of 1,000", table_cell_style)],
        [Paragraph("<b>970</b>", table_cell_style), Paragraph("97.0%", table_cell_style), Paragraph("970 parts silver out of 1,000", table_cell_style)],
        [Paragraph("<b>958</b>", table_cell_style), Paragraph("95.8%", table_cell_style), Paragraph("958 parts silver out of 1,000", table_cell_style)],
        [Paragraph("<b>925</b>", table_cell_style), Paragraph("92.5%", table_cell_style), Paragraph("925 parts silver out of 1,000 (Sterling)", table_cell_style)],
        [Paragraph("<b>835</b>", table_cell_style), Paragraph("83.5%", table_cell_style), Paragraph("835 parts silver out of 1,000", table_cell_style)],
        [Paragraph("<b>800</b>", table_cell_style), Paragraph("80.0%", table_cell_style), Paragraph("800 parts silver out of 1,000", table_cell_style)],
    ]
    t = Table(table_rows, colWidths=[90, 95, 330])
    t.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#1A1A1A')),
        ('BOTTOMPADDING', (0,0), (-1,0), 6),
        ('TOPPADDING', (0,0), (-1,0), 6),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.HexColor('#FFFFFF'), colors.HexColor('#F8F5EE')]),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#D4AF37')),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('TOPPADDING', (0,1), (-1,-1), 4),
        ('BOTTOMPADDING', (0,1), (-1,-1), 4),
    ]))

    skip_next_cheat_sheet_lines = 0

    for idx, p in enumerate(paras):
        if "SILVER HOUSE" in p or "THE SIMPLE GUIDE TO SILVER PURITY" in p or "No chemistry degree required" in p:
            continue
        
        # When reaching cheat sheet table items
        if p == "The Silver Purity Cheat Sheet":
            story.append(Spacer(1, 4))
            story.append(Paragraph("<b>The Silver Purity Cheat Sheet</b>", h1_style))
            story.append(HRFlowable(width="100%", thickness=0.8, color=colors.HexColor('#E2D7BE'), spaceBefore=2, spaceAfter=6))
            story.append(t)
            story.append(Spacer(1, 8))
            skip_next_cheat_sheet_lines = 24 # skip individual table rows in raw text
            continue

        if skip_next_cheat_sheet_lines > 0:
            if "A useful memory trick" in p:
                skip_next_cheat_sheet_lines = 0
            else:
                skip_next_cheat_sheet_lines -= 1
                continue

        # Major Headings
        if p in [
            "First: What Does Purity Mean?",
            "999 Silver: Almost Pure Silver",
            "925 Silver: Sterling Silver",
            "Why Not Make Everything 100% Silver?",
            "Weight and Purity Are Different Things",
            "The One Formula Worth Remembering",
            "Now Comes the Important Word: Hallmark",
            "Reading the Silver Hallmark",
            "Can You Verify a Hallmark?",
            "One More Important Point: Purity Is Not the Whole Story",
            "The Silver House Way",
            "The 10-Second Memory Trick"
        ]:
            story.append(Spacer(1, 4))
            story.append(Paragraph(f"<b>{p}</b>", h1_style))
            story.append(HRFlowable(width="100%", thickness=0.6, color=colors.HexColor('#E2D7BE'), spaceBefore=2, spaceAfter=4))
        
        # Formula Box
        elif "Pure Silver Content = Weight" in p:
            f_box = [[
                Paragraph("<b>Pure Silver Content = Weight &times; Purity &divide; 1000</b>", box_style)
            ]]
            ft = Table(f_box, colWidths=[515])
            ft.setStyle(TableStyle([
                ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#F8F5EE')),
                ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#D4AF37')),
                ('ALIGN', (0,0), (-1,-1), 'CENTER'),
                ('TOPPADDING', (0,0), (-1,-1), 6),
                ('BOTTOMPADDING', (0,0), (-1,-1), 6),
            ]))
            story.append(Spacer(1, 4))
            story.append(ft)
            story.append(Spacer(1, 4))

        # Memory Trick Summary Box
        elif p.startswith("• 999") or p.startswith("• 925") or p.startswith("• Purity") or p.startswith("• Weight") or p.startswith("• Hallmark"):
            story.append(Paragraph(f"<b>&bull;</b> {p[2:]}", body_style))

        elif "Silver House | Know your silver" in p:
            end_box = [[
                Paragraph(f"<b>{p}</b>", ParagraphStyle('EndP', parent=box_style, alignment=1, textColor=colors.HexColor('#9E7D3B')))
            ]]
            et = Table(end_box, colWidths=[515])
            et.setStyle(TableStyle([
                ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#FBF9F5')),
                ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#D4AF37')),
                ('TOPPADDING', (0,0), (-1,-1), 8),
                ('BOTTOMPADDING', (0,0), (-1,-1), 8),
            ]))
            story.append(Spacer(1, 8))
            story.append(et)

        else:
            story.append(Paragraph(p, body_style))

    doc.build(story, canvasmaker=NumberedCanvas)
    print(f"[OK] Built {output_pdf} ({os.path.getsize(output_pdf)} bytes)")

if __name__ == '__main__':
    logo = 'public/images/silver_house_seal.jpg'
    build_about_us_pdf('Silver_House_About_Us_Full.docx', 'public/docs/Silver_House_About_Us.pdf', logo)
    build_purity_guide_pdf('Silver_House_Silver_Purity_Guide.docx', 'public/docs/Silver_House_Silver_Purity_Guide.pdf', logo)
